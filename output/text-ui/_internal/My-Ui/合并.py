import os
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from copy import copy
import argparse

def copy_cell_style(src_cell, dst_cell):
    """复制单元格样式"""
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.border = copy(src_cell.border)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy(src_cell.protection)
        dst_cell.alignment = copy(src_cell.alignment)

def copy_sheet(src_sheet, dst_sheet):
    """复制整个工作表，包括格式和公式"""
    # 复制单元格值和样式
    for row in src_sheet.iter_rows():
        for cell in row:
            if cell.value is not None or cell.has_style:
                dst_cell = dst_sheet.cell(row=cell.row, column=cell.column, value=cell.value)
                copy_cell_style(cell, dst_cell)
    
    # 复制列宽
    for col in range(1, src_sheet.max_column + 1):
        col_letter = get_column_letter(col)
        if col_letter in src_sheet.column_dimensions:
            dst_sheet.column_dimensions[col_letter].width = src_sheet.column_dimensions[col_letter].width
    
    # 复制行高
    for row in range(1, src_sheet.max_row + 1):
        if row in src_sheet.row_dimensions:
            dst_sheet.row_dimensions[row].height = src_sheet.row_dimensions[row].height
    
    # 复制合并的单元格
    for merged_range in src_sheet.merged_cells.ranges:
        dst_sheet.merge_cells(str(merged_range))

def merge_excel_files(output_filename="合并后的Excel工作簿.xlsx", include_hidden=False):
    """
    合并当前文件夹中的所有Excel文件到一个工作簿中
    
    Args:
        output_filename (str): 输出文件名
        include_hidden (bool): 是否包含隐藏的工作表
    """
    # 获取当前目录
    current_dir = os.getcwd()
    
    # 查找所有Excel文件
    excel_files = [f for f in os.listdir(current_dir) 
                  if f.endswith(('.xlsx', '.xls')) and not f.startswith('~$')]
    
    if not excel_files:
        print("当前文件夹中没有找到Excel文件!")
        return False
    
    print(f"找到 {len(excel_files)} 个Excel文件:")
    for i, file in enumerate(excel_files, 1):
        print(f"  {i}. {file}")
    
    # 创建新的工作簿
    new_wb = Workbook()
    # 删除默认创建的工作表
    new_wb.remove(new_wb.active)
    
    output_file = os.path.join(current_dir, output_filename)
    
    print(f"\n开始合并文件到: {output_filename}")
    
    success_count = 0
    error_count = 0
    
    for i, file in enumerate(excel_files, 1):
        print(f"正在处理 ({i}/{len(excel_files)}): {file}")
        
        try:
            # 读取Excel文件
            file_path = os.path.join(current_dir, file)
            
            # 使用openpyxl读取文件以保留格式
            src_wb = load_workbook(file_path, data_only=False)
            
            # 处理每个工作表
            for sheet_name in src_wb.sheetnames:
                # 跳过隐藏的工作表（如果选项未选中）
                if src_wb[sheet_name].sheet_state == 'hidden' and not include_hidden:
                    continue
                    
                src_sheet = src_wb[sheet_name]
                
                # 创建工作表名称（使用原始文件名）
                new_sheet_name = os.path.splitext(file)[0]
                if len(new_sheet_name) > 31:
                    new_sheet_name = new_sheet_name[:31]
                
                # 处理名称重复的情况
                original_name = new_sheet_name
                counter = 1
                while new_sheet_name in new_wb.sheetnames:
                    new_sheet_name = f"{original_name}_{counter}"
                    if len(new_sheet_name) > 31:
                        new_sheet_name = new_sheet_name[:31]
                    counter += 1
                
                # 创建新工作表
                new_sheet = new_wb.create_sheet(title=new_sheet_name)
                
                # 复制工作表内容（包括格式）
                copy_sheet(src_sheet, new_sheet)
                print(f"  ✓ 已添加工作表: {new_sheet_name}")
            
            src_wb.close()
            success_count += 1
            
        except Exception as e:
            print(f"  ✗ 处理文件 {file} 时出错: {str(e)}")
            error_count += 1
            continue
    
    # 保存新工作簿
    print(f"\n正在保存工作簿...")
    new_wb.save(output_file)
    new_wb.close()
    
    print(f"\n合并完成!")
    print(f"成功处理: {success_count} 个文件")
    print(f"处理失败: {error_count} 个文件")
    print(f"输出文件: {output_file}")
    
    return True

def main():
    """主函数"""
    parser = argparse.ArgumentParser(description='合并Excel文件并保留格式')
    parser.add_argument('-o', '--output', default='合并后的Excel工作簿.xlsx', 
                       help='输出文件名 (默认: 合并后的Excel工作簿.xlsx)')
    parser.add_argument('--include-hidden', action='store_true',
                       help='包含隐藏的工作表')
    
    args = parser.parse_args()
    
    print("=" * 50)
    print("Excel 文件合并工具 (保留格式)")
    print("=" * 50)
    
    merge_excel_files(args.output, args.include_hidden)

if __name__ == "__main__":
    main()