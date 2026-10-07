import sys
import pandas as pd
import openpyxl
from PyQt5 import QtCore, QtGui, QtWidgets
import 表三  # 自定义模块，处理表三生成
import 表四  # 自定义模块，处理表四生成
import 单盘测试  # 自定义模块，处理单盘测试


class Ui_Form:
    """主界面类，用于生成测试表格"""
    
    def __init__(self):
        super().__init__()
        self.filePaths = []  # 存储文件路径的列表

    def setupUi(self, Form):
        """设置UI界面"""
        Form.setObjectName("Form")
        Form.resize(807, 659)
        
        # 创建UI控件
        self._create_controls(Form)
        
        # 设置控件属性
        self._setup_controls()
        
        # 连接信号和槽
        self._connect_signals()
        
        self.retranslateUi(Form)
        QtCore.QMetaObject.connectSlotsByName(Form)

    def _create_controls(self, Form):
        """创建界面控件"""
        # 单盘测试区域
        self.groupBox = QtWidgets.QGroupBox(Form)
        self.groupBox.setGeometry(QtCore.QRect(40, 40, 551, 121))
        
        # 光缆类型选择
        self.groupBox_2 = QtWidgets.QGroupBox(self.groupBox)
        self.groupBox_2.setGeometry(QtCore.QRect(20, 17, 121, 91))
        self.radioButton = QtWidgets.QRadioButton(self.groupBox_2)  # 普通光缆
        self.radioButton_2 = QtWidgets.QRadioButton(self.groupBox_2)  # 防鼠光缆
        
        # 操作按钮
        self.pushButton = QtWidgets.QPushButton(self.groupBox)  # 开始创建
        self.pushButton_2 = QtWidgets.QPushButton(self.groupBox)  # 合并表格
        
        # 光缆型号选择
        self.groupBox_3 = QtWidgets.QGroupBox(self.groupBox)
        self.groupBox_3.setGeometry(QtCore.QRect(160, 17, 120, 91))
        self.comboBox = QtWidgets.QComboBox(self.groupBox_3)  # 型号下拉框
        
        # 打印数量输入
        self.groupBox_4 = QtWidgets.QGroupBox(self.groupBox)
        self.groupBox_4.setGeometry(QtCore.QRect(300, 17, 121, 91))
        self.textEdit = QtWidgets.QTextEdit(self.groupBox_4)  # 数量输入框
        
        # 标签
        self.label = QtWidgets.QLabel(Form)
        self.label_2 = QtWidgets.QLabel(Form)
        self.label_3 = QtWidgets.QLabel(Form)
        self.label_4 = QtWidgets.QLabel(Form)
        self.label_5 = QtWidgets.QLabel(Form)
        self.label_6 = QtWidgets.QLabel(Form)
        
        # 文件选择区域
        self.entry1_text = QtWidgets.QLineEdit(Form)  # 文件路径显示
        self.pushButton_3 = QtWidgets.QPushButton(Form)  # 选择路径按钮
        self.pushButton_4 = QtWidgets.QPushButton(Form)  # 制作表三按钮
        self.pushButton_5 = QtWidgets.QPushButton(Form)  # 制作表四按钮
        
        # 结果显示表格
        self.tableWidget = QtWidgets.QTableWidget(Form)
        self.tableWidget.setGeometry(QtCore.QRect(40, 250, 721, 371))

    def _setup_controls(self):
        """设置控件属性"""
        # 设置单选按钮默认状态
        self.radioButton.setChecked(True)
        self.radioButton.setGeometry(QtCore.QRect(20, 30, 89, 16))
        self.radioButton_2.setGeometry(QtCore.QRect(20, 60, 89, 16))
        
        # 设置按钮位置
        self.pushButton.setGeometry(QtCore.QRect(450, 27, 75, 31))
        self.pushButton_2.setGeometry(QtCore.QRect(450, 75, 75, 31))
        
        # 设置下拉框选项
        self.comboBox.setGeometry(QtCore.QRect(20, 31, 81, 31))
        self.comboBox.addItems(["4D", "8D", "12D", "24D", "48D", "96D", "144D"])
        
        # 设置文本输入框
        self.textEdit.setGeometry(QtCore.QRect(20, 31, 81, 31))
        
        # 设置文件路径输入框为只读
        self.entry1_text.setReadOnly(True)
        self.entry1_text.setGeometry(QtCore.QRect(120, 180, 291, 31))
        
        # 设置按钮位置
        self.pushButton_3.setGeometry(QtCore.QRect(420, 180, 101, 31))
        self.pushButton_4.setGeometry(QtCore.QRect(530, 180, 101, 31))
        self.pushButton_5.setGeometry(QtCore.QRect(640, 180, 101, 31))

    def _connect_signals(self):
        """连接信号和槽函数"""
        self.pushButton.clicked.connect(self.on_pushButton_clicked)
        self.pushButton_3.clicked.connect(self.get_path)
        self.pushButton_4.clicked.connect(self.df_table_1)
        self.pushButton_5.clicked.connect(self.glcs_b3)

    def retranslateUi(self, Form):
        """设置界面文本翻译"""
        _translate = QtCore.QCoreApplication.translate
        Form.setWindowTitle(_translate("Form", "制作测试表格"))
        self.groupBox.setTitle(_translate("Form", "单盘测试"))
        self.groupBox_2.setTitle(_translate("Form", "光缆类型："))
        self.radioButton.setText(_translate("Form", "普通光缆"))
        self.radioButton_2.setText(_translate("Form", "防鼠光缆"))
        self.pushButton.setText(_translate("Form", "开始创建"))
        self.groupBox_3.setTitle(_translate("Form", "光缆型号："))
        self.groupBox_4.setTitle(_translate("Form", "打印数量"))
        self.pushButton_2.setText(_translate("Form", "合并表格"))
        self.label.setText(_translate("Form", "TextLabel1"))
        self.label_2.setText(_translate("Form", "TextLabel2"))
        self.label_3.setText(_translate("Form", "TextLabel3"))
        self.label_4.setText(_translate("Form", "TextLabel"))
        self.label_5.setText(_translate("Form", "TextLabel"))
        self.label_6.setText(_translate("Form", "选择文件："))
        self.pushButton_3.setText(_translate("Form", "选择路径"))
        self.pushButton_4.setText(_translate("Form", "开始制作表三"))
        self.pushButton_5.setText(_translate("Form", "开始制作表四"))

    def on_pushButton_clicked(self):
        """处理开始创建按钮点击事件"""
        # 获取光缆类型
        型号 = "GTYA53" if self.radioButton_2.isChecked() else "GTYA"
        
        # 获取光缆芯数
        光缆芯数 = int(self.comboBox.currentText()[:-1])
        
        # 获取打印页数
        try:
            页数 = int(self.textEdit.toPlainText())
        except ValueError:
            QtWidgets.QMessageBox.warning(self.pushButton, "警告", "请输入有效的数字！")
            return
        
        # 执行单盘测试
        try:
            单盘测试.dpcs_zcx(型号, 光缆芯数, 页数)
            QtWidgets.QMessageBox.information(self.pushButton, "提示", "单盘测试完成！")
        except Exception as e:
            QtWidgets.QMessageBox.critical(self.pushButton, "错误", f"执行单盘测试时出错：{str(e)}")

    def get_path(self):
        """获取文件路径"""
        path1, _ = QtWidgets.QFileDialog.getOpenFileName(
            self.pushButton_3, "选择文件", "", "Excel Files (*.xlsx)"
        )
        if path1:
            self.entry1_text.setText(path1)
            self.filePaths.append(path1)

    def _read_excel_data(self):
        """读取Excel数据"""
        if not self.filePaths:
            QtWidgets.QMessageBox.warning(None, "警告", "请先选择文件！")
            return None
        
        try:
            df = pd.read_excel(self.filePaths[0], sheet_name='数据')
            df = df.iloc[:, [1, 2, 3, 4]]  # 选择需要的列
            df = df.dropna(subset=['接头数'])  # 去除接头数为空的行
            return df
        except Exception as e:
            QtWidgets.QMessageBox.critical(None, "错误", f"读取Excel文件时出错：{str(e)}")
            return None

    def df_table_1(self):
        """生成表三"""
        df = self._read_excel_data()
        if df is None:
            return
        
        # 创建工作簿
        wb = openpyxl.Workbook()
        sheet1 = wb.active
        
        # 设置列宽
        for col in range(1, 11):
            sheet1.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 8.4
        
        # 处理每一行数据
        for i in range(len(df)):
            地址 = df.iloc[i, 0]
            光缆芯数 = int(df.iloc[i, 1][:-1])
            光缆总长 = df.iloc[i, 2]
            接头数 = int(df.iloc[i, 3])
            
            表三.表三(sheet1, 光缆芯数, 光缆总长, 接头数, 地址)
        
        # 保存文件
        try:
            wb.save("表三接头损耗表.xlsx")
            QtWidgets.QMessageBox.information(None, "提示", "表三制作完成！")
        except Exception as e:
            QtWidgets.QMessageBox.critical(None, "错误", f"保存文件时出错：{str(e)}")

    def glcs_b3(self):
        """生成表四"""
        df = self._read_excel_data()
        if df is None:
            return
        
        # 创建工作簿
        wb = openpyxl.Workbook()
        sheet1 = wb.active
        
        # 设置列宽
        for col in range(1, 13):
            sheet1.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 7.28
        
        # 处理每一行数据
        for i in range(len(df)):
            地址 = df.iloc[i, 0]
            光缆芯数 = int(df.iloc[i, 1][:-1])
            formatted_length = df.iloc[i, 2] / 1000
            光缆总长 = f"{formatted_length:.4f}"
            
            表四.glcs_b3(sheet1, 光缆芯数, 光缆总长, 地址)
        
        # 保存文件
        try:
            wb.save("表四损耗表.xlsx")
            QtWidgets.QMessageBox.information(None, "提示", "表四制作完成！")
        except Exception as e:
            QtWidgets.QMessageBox.critical(None, "错误", f"保存文件时出错：{str(e)}")


if __name__ == "__main__":
    app = QtWidgets.QApplication(sys.argv)
    form = QtWidgets.QWidget()
    ui = Ui_Form()
    ui.setupUi(form)
    form.show()
    sys.exit(app.exec_())