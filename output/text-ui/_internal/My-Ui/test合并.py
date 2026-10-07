import os
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from copy import copy
import argparse
from io import StringIO
import sys
from datetime import datetime


class ExcelMergerLogger:
    """日志收集器类"""
    def __init__(self):
        self.log_buffer = StringIO()
        self.log_lines = []
    
    def log(self, message):
        """添加日志消息"""
        timestamp = datetime.now().strftime("%H:%M:%S")
        log_message = f"[{timestamp}] {message}"
        self.log_lines.append(log_message)
        self.log_buffer.write(log_message + "\n")
        print(log_message)  # 同时输出到控制台
        return log_message
    
    def get_logs(self):
        """获取所有日志内容"""
        return self.log_buffer.getvalue()
    
    def get_last_log(self):
        """获取最后一条日志"""
        return self.log_lines[-1] if self.log_lines else ""
    
    def clear_logs(self):
        """清空日志"""
        self.log_buffer = StringIO()
        self.log_lines = []

# 创建全局日志器实例
logger = ExcelMergerLogger()

def copy_cell_style(src_cell, dst_cell):
    """复制单元格样式"""
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.border = copy(src_cell.border)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy(src_cell.protection)
        dst_cell.alignment = copy(src_cell.alignment)

def copy_sheet(src_sheet, dst_sheet):
    """复制整个工作表，包括格式和公式"""
    # 复制单元格值和样式
    for row in src_sheet.iter_rows():
        for cell in row:
            if cell.value is not None or cell.has_style:
                dst_cell = dst_sheet.cell(row=cell.row, column=cell.column, value=cell.value)
                copy_cell_style(cell, dst_cell)
    
    # 复制列宽
    for col in range(1, src_sheet.max_column + 1):
        col_letter = get_column_letter(col)
        if col_letter in src_sheet.column_dimensions:
            dst_sheet.column_dimensions[col_letter].width = src_sheet.column_dimensions[col_letter].width
    
    # 复制行高
    for row in range(1, src_sheet.max_row + 1):
        if row in src_sheet.row_dimensions:
            dst_sheet.row_dimensions[row].height = src_sheet.row_dimensions[row].height
    
    # 复制合并的单元格
    for merged_range in src_sheet.merged_cells.ranges:
        dst_sheet.merge_cells(str(merged_range))

def merge_excel_files(file_list=None, output_filename="合并后的Excel工作簿.xlsx", 
                     output_dir=None, include_hidden=False, callback=None):
    """
    合并指定的Excel文件到一个工作簿中
    
    Args:
        file_list (list): 要合并的Excel文件路径列表
        output_filename (str): 输出文件名
        output_dir (str): 输出目录，默认为当前目录或第一个文件所在目录
        include_hidden (bool): 是否包含隐藏的工作表
        callback (function): 回调函数，用于更新UI进度
                        格式: callback(status, current, total, percentage, message)
    
    Returns:
        dict: 包含结果信息的字典
    """
    from datetime import datetime
    
    # 清空之前的日志
    logger.clear_logs()
    
    # 处理文件列表
    if file_list is None:
        # 如果没有提供文件列表，使用当前目录中的所有Excel文件
        current_dir = os.getcwd()
        excel_files = [f for f in os.listdir(current_dir) 
                      if f.endswith(('.xlsx', '.xls')) and not f.startswith('~$')]
        file_list = [os.path.join(current_dir, f) for f in excel_files]
    else:
        # 验证文件列表中的文件是否存在
        valid_files = []
        for file_path in file_list:
            if os.path.exists(file_path):
                valid_files.append(file_path)
            else:
                logger.log(f"警告: 文件不存在，已跳过: {file_path}")
        file_list = valid_files
    
    if not file_list:
        error_msg = "没有找到有效的Excel文件!"
        logger.log(error_msg)
        if callback:
            callback('error', 0, 0, 0, error_msg)
        return {
            'success': False,
            'message': error_msg,
            'output_file': None,
            'processed_count': 0,
            'error_count': 0,
            'total_count': 0
        }
    
    # 确定输出目录
    if output_dir is None:
        if file_list:
            output_dir = os.path.dirname(file_list[0])
        else:
            output_dir = os.getcwd()
    
    # 确保输出目录存在
    os.makedirs(output_dir, exist_ok=True)
    output_file = os.path.join(output_dir, output_filename)
    
    # 记录找到的文件
    logger.log(f"找到 {len(file_list)} 个Excel文件")
    for i, file_path in enumerate(file_list, 1):
        logger.log(f"  {i}. {os.path.basename(file_path)}")
    
    # 创建新的工作簿
    new_wb = Workbook()
    # 删除默认创建的工作表
    new_wb.remove(new_wb.active)
    
    logger.log(f"开始合并文件到: {output_filename}")
    
    success_count = 0
    error_count = 0
    
    # 如果有回调函数，通知UI总文件数
    if callback:
        callback('init', 0, len(file_list), 0, f"开始处理 {len(file_list)} 个文件")
    
    # 处理每个文件
    for i, file_path in enumerate(file_list, 1):
        file_name = os.path.basename(file_path)
        
        # 计算进度百分比
        percentage = int((i - 1) / len(file_list) * 100)
        progress_msg = f"正在处理 ({i}/{len(file_list)}): {file_name}"
        logger.log(progress_msg)
        
        # 更新UI进度
        if callback:
            callback('progress', i, len(file_list), percentage, progress_msg)
        
        try:
            # 读取Excel文件
            src_wb = load_workbook(file_path, data_only=False)
            
            # 获取工作表列表
            sheet_names = [name for name in src_wb.sheetnames 
                          if include_hidden or src_wb[name].sheet_state != 'hidden']
            
            # 处理每个工作表
            for sheet_index, sheet_name in enumerate(sheet_names, 1):
                src_sheet = src_wb[sheet_name]
                
                # 创建工作表名称（使用原始文件名）
                new_sheet_name = os.path.splitext(file_name)[0]
                if len(new_sheet_name) > 31:
                    new_sheet_name = new_sheet_name[:31]
                
                # 处理名称重复的情况
                original_name = new_sheet_name
                counter = 1
                while new_sheet_name in new_wb.sheetnames:
                    new_sheet_name = f"{original_name}_{counter}"
                    if len(new_sheet_name) > 31:
                        new_sheet_name = new_sheet_name[:31]
                    counter += 1
                
                # 创建新工作表
                new_sheet = new_wb.create_sheet(title=new_sheet_name)
                
                # 复制工作表内容（包括格式）
                copy_sheet(src_sheet, new_sheet)
                sheet_msg = f"  已添加工作表: {new_sheet_name}"
                logger.log(sheet_msg)
                
                # 工作表级别的进度回调（可选）
                if callback and len(sheet_names) > 1:
                    sheet_percentage = percentage + int(sheet_index / len(sheet_names) * (100 / len(file_list)))
                    callback('sheet_progress', sheet_index, len(sheet_names), sheet_percentage, sheet_msg)
            
            src_wb.close()
            success_count += 1
            
            # 文件处理完成回调
            if callback:
                file_percentage = int(i / len(file_list) * 100)
                callback('file_complete', i, len(file_list), file_percentage, f"完成处理: {file_name}")
            
        except Exception as e:
            error_msg = f"处理文件 {file_name} 时出错: {str(e)}"
            logger.log(error_msg)
            error_count += 1
            
            # 错误回调
            if callback:
                callback('error', i, len(file_list), percentage, error_msg)
            continue
    
    # 保存新工作簿
    save_msg = "正在保存工作簿..."
    logger.log(save_msg)
    if callback:
        callback('saving', 0, 0, 95, save_msg)
    
    try:
        new_wb.save(output_file)
        new_wb.close()
        
        # 完成消息
        completion_msg = f"合并完成! 成功: {success_count}, 失败: {error_count}"
        logger.log(completion_msg)
        
        if callback:
            callback('complete', success_count, error_count, 100, completion_msg)
        
        return {
            'success': True,
            'message': completion_msg,
            'output_file': output_file,
            'processed_count': success_count + error_count,
            'success_count': success_count,
            'error_count': error_count,
            'total_count': len(file_list),
            'logs': logger.get_logs()
        }
    
    except Exception as e:
        error_msg = f"保存文件时出错: {str(e)}"
        logger.log(error_msg)
        
        if callback:
            callback('error', 0, 0, 95, error_msg)
        
        return {
            'success': False,
            'message': error_msg,
            'output_file': None,
            'processed_count': success_count + error_count,
            'success_count': success_count,
            'error_count': error_count + 1,
            'total_count': len(file_list),
            'logs': logger.get_logs()
        }

# 独立使用时的主函数
def main():
    """主函数（独立使用时的入口点）"""
    parser = argparse.ArgumentParser(description='合并Excel文件并保留格式')
    parser.add_argument('-o', '--output', default='合并后的Excel工作簿.xlsx', 
                       help='输出文件名 (默认: 合并后的Excel工作簿.xlsx)')
    parser.add_argument('-d', '--output-dir', 
                       help='输出目录 (默认: 第一个文件所在目录或当前目录)')
    parser.add_argument('--include-hidden', action='store_true',
                       help='包含隐藏的工作表')
    parser.add_argument('files', nargs='*', help='要合并的Excel文件路径')
    
    args = parser.parse_args()
    
    print("=" * 50)
    print("Excel 文件合并工具 (保留格式)")
    print("=" * 50)
    
    # 简单的控制台回调函数
    def console_callback(status, current, total, percentage, message=None):
        if status == 'progress' and message:
            print(f"[{percentage}%] {message}")
        elif status == 'complete':
            print(f"[100%] {message}")
        elif status == 'error':
            print(f"[错误] {message}")
    
    result = merge_excel_files(
        file_list=args.files if args.files else None,
        output_filename=args.output,
        output_dir=args.output_dir,
        include_hidden=args.include_hidden,
        callback=console_callback
    )
    
    if result['success']:
        print(f"\n{result['message']}")
        print(f"输出文件: {result['output_file']}")
    else:
        print(f"\n错误: {result['message']}")

# 供其他程序调用的接口函数
def start_merge(file_list, output_filename="合并后的Excel工作簿.xlsx", 
               output_dir=None, include_hidden=False, progress_callback=None):
    """
    从其他程序启动合并过程的主要函数
    
    Args:
        file_list (list): 要合并的Excel文件路径列表
        output_filename (str): 输出文件名
        output_dir (str): 输出目录
        include_hidden (bool): 是否包含隐藏的工作表
        progress_callback (function): 进度回调函数
                        格式: callback(status, current, total, percentage, message)
        
    Returns:
        dict: 包含结果信息的字典
    """
    return merge_excel_files(
        file_list=file_list,
        output_filename=output_filename,
        output_dir=output_dir,
        include_hidden=include_hidden,
        callback=progress_callback
    )

def start_merge_from_folder(folder_path, output_filename="合并后的Excel工作簿.xlsx", 
                           output_dir=None, include_hidden=False, progress_callback=None):
    """
    合并指定文件夹中的所有Excel文件
    
    Args:
        folder_path (str): 文件夹路径
        output_filename (str): 输出文件名
        output_dir (str): 输出目录
        include_hidden (bool): 是否包含隐藏的工作表
        progress_callback (function): 进度回调函数
                        格式: callback(status, current, total, percentage, message)
        
    Returns:
        dict: 包含结果信息的字典
    """
    # 获取文件夹中的所有Excel文件
    if not os.path.exists(folder_path):
        error_msg = f"文件夹不存在: {folder_path}"
        if progress_callback:
            progress_callback('error', 0, 0, 0, error_msg)
        return {
            'success': False,
            'message': error_msg,
            'output_file': None
        }
    
    excel_files = [os.path.join(folder_path, f) for f in os.listdir(folder_path) 
                  if f.endswith(('.xlsx', '.xls')) and not f.startswith('~$')]
    
    return merge_excel_files(
        file_list=excel_files,
        output_filename=output_filename,
        output_dir=output_dir,
        include_hidden=include_hidden,
        callback=progress_callback
    )

def get_logs():
    """获取当前日志内容"""
    return logger.get_logs()

def get_last_log():
    """获取最后一条日志"""
    return logger.get_last_log()

def clear_logs():
    """清空日志"""
    logger.clear_logs()

def get_logger_instance():
    """获取日志器实例，供其他程序直接操作"""
    return logger

